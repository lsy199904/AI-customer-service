"""Rebuild the curated Word source. Authoring requires python-docx, runtime ingestion does not."""
from pathlib import Path
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parent.parent

def build():
 items=[
 ('SmartTITU 可以帮助企业做什么','SmartTITU 面向企业提供模型聚合与 Token 运营服务。TITU Hub 负责模型聚合与运营，TITU Edge 负责企业 Token 管控，TITU Work 定位为企业智能体工作平台。具体功能与开通范围以客户当前版本和正式交付为准。','14.SmartTITU智信词元产品体系介绍-20261020.pdf，物理第2页。'),
 ('如何注册和登录企业账号','在登录页面填写手机号和验证码，提交登录。系统根据过往手机号信息判断是否已注册。新客户会看到注册页面，按页面要求填写企业注册信息并提交，系统保存资料、完成注册后进入工作台。已经注册的客户登录成功后直接进入工作台。进入工作台后，尚未认证的企业需要继续办理企业认证，认证完成后方可领取试用套餐及订购正式套餐。','flowchart-transposed.html，企业客户注册、登录及企业认证业务流程图。'),
 ('企业注册需要填写哪些信息','新客户在注册页面按要求填写企业名称、组织机构代码证、行业、行业分类、推荐客户经理和公司地址等企业信息。填写完成后提交，系统保存企业客户资料并完成注册。具体填写要求按当前注册页面指引操作。','flowchart-transposed.html，企业客户注册、登录及企业认证业务流程图。'),
 ('企业认证如何办理','登录工作台后，系统判断企业是否已完成认证；未认证时会提示继续办理。企业可选择在线或线下认证：在线认证按页面提供的法大大或 e签宝指引完成；线下认证联系客户经理或致电热线，由客户经理受理并协助完成。企业认证完成后，方可领取试用套餐及订购正式套餐。','flowchart-transposed.html，企业客户注册、登录及企业认证业务流程图。'),
 ('如何创建 API 密钥','在企业服务门户左侧进入“服务 → API密钥”，创建并设置密钥名称。可以配置 IP 调用白名单和允许调用的模型。创建成功后及时保存，密钥仅显示一次，不要通过聊天或公开文档发送完整密钥。','11.TITU Hub用户操作手册-20261009.pdf，物理第5页。'),
 ('在哪里查找模型 ID','在模型广场选择已开通模型，查看并复制实际 model ID。调用请求中的 model 字段使用该 ID。可用模型、上下文长度和价格以当前模型广场展示为准。','11.TITU Hub用户操作手册-20261009.pdf，物理第4页。'),
 ('模型接口的基础地址是什么','模型 API 基础地址为 https://api.smarttitu.cn/v1/。Chat Completions 请求路径为 /chat/completions。请求通过 Authorization 头传入 Bearer 加实际 API 密钥，并使用已开通的 model ID。图片和视频模型的接口能力可能不同，应按对应调用示例使用。','7.SmartTITU销售常见问题速查FAQ-20261008v1.pdf，物理第7页。'),
 ('平台费用如何计算','资料说明平台按模型单价和实际 Token 用量计费，支持预付和后付费。模型价格、折扣、账期与结算条款请以当前控制台、书面报价和合同为准，不使用历史资料中的估算作为承诺。','7.SmartTITU销售常见问题速查FAQ-20261008v1.pdf，物理第3页。'),
 ('怎样查看调用用量和明细','在“服务 → 用量统计”查看调用趋势和明细。可以按 API 密钥、模型、成功或失败状态以及日期范围筛选，并按页面能力导出记录用于核对。','11.TITU Hub用户操作手册-20261009.pdf，物理第6页。'),
 ('忘记登录密码或企业信息需要修改怎么办','企业管理页面可查看企业信息。如需重置登录密码，或修改企业名称、地址等信息，请联系专属客户经理处理。','7.SmartTITU销售常见问题速查FAQ-20261008v1.pdf，物理第2至3页。'),
 ('模型调用失败时先检查哪些信息','先记录失败时间、模型 ID、密钥名称、状态码和 Request ID。检查请求地址、鉴权头、密钥权限、来源 IP 及网络。不要在反馈中提供完整密钥。根据状态码继续排查；没有状态码时先获取完整报错内容，不能把调用失败直接解释为计费问题。','7.SmartTITU销售常见问题速查FAQ-20261008v1.pdf，物理第8至9页。'),
 ('400 请求错误如何排查','400 通常表示请求参数错误或模型能力不支持。检查请求体结构、必填参数和 model ID；如果使用函数或工具调用，确认模型支持对应能力。','7.SmartTITU销售常见问题速查FAQ-20261008v1.pdf，物理第8页。'),
 ('401 鉴权失败如何排查','401 表示鉴权失败。检查请求头是否为 Authorization: Bearer 加实际密钥，确认密钥有效，并核对模型权限与 IP 白名单。不要将密钥名称误当成完整密钥。','7.SmartTITU销售常见问题速查FAQ-20261008v1.pdf，物理第8页。'),
 ('403 权限不足如何排查','403 表示当前密钥没有相应模型或接口的访问权限，或请求来源 IP 不在白名单。核对目标模型授权和实际出口 IP。','7.SmartTITU销售常见问题速查FAQ-20261008v1.pdf，物理第8页。'),
 ('404 资源不存在如何排查','404 可能与 model ID 或请求路径错误有关。对照模型广场中的实际模型 ID 和对应接口示例，检查基础地址和请求路径是否重复或缺失。','7.SmartTITU销售常见问题速查FAQ-20261008v1.pdf，物理第8页。'),
 ('429 限流如何处理','429 通常表示触发每分钟请求数 RPM 或每分钟 Token 数 TPM 限制。降低请求频率或错峰调用；容量需求较大时联系客户经理申请调整。实际限额以账号及密钥配置为准。','7.SmartTITU销售常见问题速查FAQ-20261008v1.pdf，物理第8页。'),
 ('500 或 503 服务异常如何处理','500 或 503 可能是平台或上游暂时异常。可以稍后重试；503 也可能表示当前模型没有可用推理渠道。持续发生时提供调用时间、模型 ID 和 Request ID 联系技术支持。','7.SmartTITU销售常见问题速查FAQ-20261008v1.pdf，物理第8至9页。'),
 ('504 或请求超时如何排查','检查本机和服务器的网络，尝试缩短输入后重试。如果持续超时，记录模型 ID、失败时间和 Request ID，联系技术支持协助定位。不要仅凭超时判断模型或平台一定故障。','7.SmartTITU销售常见问题速查FAQ-20261008v1.pdf，物理第8至9页。'),
 ('如何联系人工客服','可联系专属客户经理，或拨打服务热线 4008841883，发送邮件至 help@china-entercom.net。反馈时提供公司与联系人信息、问题描述、调用时间、模型 ID 和 Request ID，不发送完整 API 密钥。','7.SmartTITU销售常见问题速查FAQ-20261008v1.pdf，物理第13至14页。')]
 doc=Document();section=doc.sections[0];section.page_width=Inches(8.27);section.page_height=Inches(11.69);section.top_margin=Inches(.65);section.bottom_margin=Inches(.65);section.left_margin=Inches(.8);section.right_margin=Inches(.8)
 for name in ('Normal','Title','Heading 1'):
  style=doc.styles[name];style.font.name='Arial Unicode MS';style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial Unicode MS');style.font.color.rgb=RGBColor(0,0,0)
 doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(3);doc.styles['Normal'].paragraph_format.line_spacing=1.08
 doc.styles['Title'].font.size=Pt(22);doc.styles['Title'].paragraph_format.space_after=Pt(8)
 doc.styles['Heading 1'].font.size=Pt(12);doc.styles['Heading 1'].paragraph_format.space_before=Pt(8);doc.styles['Heading 1'].paragraph_format.space_after=Pt(3)
 doc.add_paragraph('SmartTITU 客户问答基础资料','Title')
 doc.add_paragraph('文档编号：smarttitu-customer-faq\n版本：2026-10-22 v2\n适用范围：客户产品咨询与基础调用排查')
 doc.add_paragraph('本文整理账号开通、模型接入、费用说明和常见错误处理，可作为首份客服知识入库文档。各问答保留原始资料来源；实时价格、权限及已开通功能以客户当前控制台和正式交付为准。')
 for title,answer,source in items:
  doc.add_paragraph(title,'Heading 1');p=doc.add_paragraph(answer);p.paragraph_format.keep_with_next=True
  if title in ('如何注册和登录企业账号','企业注册需要填写哪些信息','企业认证如何办理'):
   p=doc.add_paragraph('配图：enterprise-register');p.paragraph_format.keep_with_next=True
   for run in p.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('616161')
  p=doc.add_paragraph('来源：'+source);p.paragraph_format.space_after=Pt(6)
  for run in p.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('616161')
 doc.add_page_break();doc.add_paragraph('企业注册与认证流程图','Heading 1')
 doc.add_picture(str(ROOT/'app/static/media/enterprise-register.png'),width=Inches(6.2))
 p=doc.add_paragraph('来源：flowchart-transposed.html，企业客户注册、登录及企业认证业务流程图。')
 for run in p.runs:run.font.size=Pt(8);run.font.color.rgb=RGBColor.from_string('616161')
 foot=section.footer.paragraphs[0];foot.alignment=2
 field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
 for border in list(doc.styles.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 for border in list(doc.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
 doc.core_properties.title='SmartTITU 客户问答基础资料';doc.core_properties.author='';doc.core_properties.subject='客户知识入库清洗版'
 dest=ROOT/'knowledge_sources/SmartTITU客户问答.docx';doc.save(dest);print(dest);return dest
if __name__=='__main__':build()
